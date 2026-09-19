"""
Генератор службових записок у форматі .docx з фірмовою шапкою ВОКЛ.
Використовується секретарем (Abacus) для швидкого створення документів
за зразком робочих документів користувача.

Потік:
  1. Секретар отримує запит (текст/голос у Telegram або з реєстру завдань)
  2. LLM формує тіло записки за шаблоном (structure з templates/sluzhbova_zapyska.md)
  3. Цей модуль формує .docx з правильною шапкою й підписом
  4. Файл зберігається у Drive та реєструється в automation_log
"""

import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), 'templates')

# Автор за замовчуванням (з ваших документів)
DEFAULT_AUTHOR_POSITION = "Заступник медичного директора з АДІТ"
DEFAULT_AUTHOR_NAME = "А.В. Белінський"


def _load_letterhead() -> list[str]:
    """Завантажити рядки фірмової шапки."""
    path = os.path.join(TEMPLATES_DIR, 'letterhead.txt')
    with open(path, encoding='utf-8') as f:
        return [line.rstrip('\n') for line in f if line.strip()]


def generate_sluzhbova(
    addressee: list[str],
    body_paragraphs: list[str],
    request_text: str,
    author_position: str = DEFAULT_AUTHOR_POSITION,
    author_name: str = DEFAULT_AUTHOR_NAME,
    output_path: str | None = None,
) -> str:
    """
    Згенерувати службову записку у .docx.

    Args:
        addressee: рядки блоку адресата (посада, установа, ПІБ у дав. відмінку)
        body_paragraphs: абзаци тексту (факти, обґрунтування)
        request_text: текст прохання ("Прошу розглянути можливість...")
        author_position: посада автора
        author_name: ПІБ автора
        output_path: куди зберегти (за замовчуванням — briefings/)

    Returns:
        Шлях до створеного .docx
    """
    doc = Document()

    # Базовий шрифт
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    # --- Шапка (по центру) ---
    for line in _load_letterhead():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.bold = True
        run.font.size = Pt(11)

    doc.add_paragraph()  # відступ

    # --- Заголовок ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("СЛУЖБОВА ЗАПИСКА")
    tr.bold = True
    tr.font.size = Pt(14)

    doc.add_paragraph()

    # --- Адресат (праворуч) ---
    for line in addressee:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run(line)

    doc.add_paragraph()

    # --- Тіло ---
    for para in body_paragraphs:
        p = doc.add_paragraph(para)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)

    # --- Прохання ---
    if request_text:
        p = doc.add_paragraph(request_text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)

    doc.add_paragraph()

    # --- Підпис ---
    p = doc.add_paragraph()
    p.add_run("З повагою,")
    p2 = doc.add_paragraph()
    p2.add_run(author_position)
    p3 = doc.add_paragraph()
    run = p3.add_run(author_name)
    run.bold = True

    # --- Збереження ---
    if output_path is None:
        out_dir = os.path.join(os.path.dirname(__file__), '..', '..', 'briefings')
        os.makedirs(out_dir, exist_ok=True)
        fname = f"sluzhbova_{datetime.now():%Y-%m-%d_%H%M}.docx"
        output_path = os.path.join(out_dir, fname)

    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    # Демонстрація — відтворення вашої записки про нестачу анестезіологів
    path = generate_sluzhbova(
        addressee=[
            "Заступнику медичного директора",
            "з контролю якості лікувального процесу",
            "КНП «Вінницька обласна лікарня ім. М.І.Пирогова»",
            "Вигонюку А.В.",
        ],
        body_paragraphs=[
            "Повідомляю Вам про актуальну проблему нестачі штатних посад "
            "лікарів-анестезіологів у ВАІТ №2 Клінічного центру анестезіології "
            "та інтенсивної терапії, яка негативно впливає на якість та "
            "безперервність медичної допомоги. Згідно зі штатним розписом, наразі "
            "відсутні ставки лікарів-анестезіологів для забезпечення роботи у СКТ "
            "кабінеті, радіологічному відділенні та кардіохірургічній операційній. "
            "Загальна потреба в додаткових ставках становить 2,0 ставки.",
        ],
        request_text=(
            "Прошу розглянути можливість додаткового укомплектування штатних "
            "посад для забезпечення повноцінної роботи зазначених підрозділів."
        ),
    )
    print(f"✅ Демо-записку створено: {path}")
